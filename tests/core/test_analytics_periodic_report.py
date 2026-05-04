"""core/analytics/periodic_report.py の単体テスト。

定期報告書（様式第5）ドラフト生成（Step 7）の挙動を検証する。
generate_plan_draft を内部で呼ぶため e-Stat キャッシュに依存する。
"""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from core.analytics.periodic_report import (
    build_periodic_report_summary,
    generate_periodic_report,
)


@pytest.mark.asyncio
async def test_generate_periodic_report_for_manufacturing() -> None:
    result = await generate_periodic_report("sample_mfg_01", report_fy=2026)
    assert result.company_id == "sample_mfg_01"
    assert "テクノパーツ製造" in result.company_name
    assert result.report_period.startswith("2026年度")
    assert "first_class_shipper" in result.eligibility_categories or \
           "second_class_shipper" in result.eligibility_categories
    assert result.compliance.summary  # サマリーが生成される
    assert "docx" in result.files
    assert Path(result.files["docx"]).exists()


@pytest.mark.asyncio
async def test_generate_periodic_report_for_chain_franchise() -> None:
    """連鎖化事業者は荷役等項目なし（handling_time_rows が空）。"""
    result = await generate_periodic_report("sample_chain_01")
    assert "chain_franchise" in result.eligibility_categories
    assert result.handling_time_rows == []
    # 計測データはあり
    assert len(result.wait_time_rows) >= 1


@pytest.mark.asyncio
async def test_invalid_company_id_raises() -> None:
    with pytest.raises(ValueError, match="企業プロフィール"):
        await generate_periodic_report("nonexistent_id")


@pytest.mark.asyncio
async def test_compliance_uses_max_minutes_from_measurements() -> None:
    """計測データ最大値が ComplianceReport の閾値判定に使われる。"""
    result = await generate_periodic_report("sample_mfg_01")
    waiting_item = next(
        i for i in result.compliance.items if "荷待ち" in i.name
    )
    expected_max = max(w.max_minutes for w in result.wait_time_rows)
    assert waiting_item.measured_value == expected_max


@pytest.mark.asyncio
async def test_progress_rows_include_target_and_simulation() -> None:
    """進捗行は数値目標 4 つ + 主要施策 3 件 = 計 7 行程度を含む。"""
    result = await generate_periodic_report("sample_mfg_01")
    assert len(result.progress_rows) >= 4  # 数値目標 4 つは最低含まれる
    statuses = {p.status_mark for p in result.progress_rows}
    assert "△ 継続対応" in statuses  # 初年度想定なので継続対応中心


@pytest.mark.asyncio
async def test_submission_checklist_attached() -> None:
    """提出実務チェックリストが必須項目を含む。"""
    result = await generate_periodic_report("sample_mfg_01")
    assert result.submission_checklist is not None
    labels = [item.label for item in result.submission_checklist.items]
    assert "事業者の名称" in labels
    assert "法人番号" in labels
    assert "指定通知書番号" in labels


@pytest.mark.asyncio
async def test_citations_include_improvement_standards() -> None:
    """改善基準告示・物流効率化法・CO2 ガイドラインが出典に含まれる。"""
    result = await generate_periodic_report("sample_mfg_01")
    labels = [c.label for c in result.citations]
    assert any("改善のための基準" in label for label in labels)
    assert any("物流効率化法" in label for label in labels)
    assert any("Ver.3.2" in label for label in labels)


@pytest.mark.asyncio
async def test_form5_docx_contains_required_sections() -> None:
    """生成された Word が様式第5 の必須セクション見出しを含む。"""
    result = await generate_periodic_report("sample_mfg_01")
    doc = Document(result.files["docx"])
    headings = [
        p.text for p in doc.paragraphs
        if p.style and p.style.name.startswith("Heading")
    ]
    text = "\n".join(headings)
    assert "Ⅰ 特定荷主の名称等" in text
    assert "Ⅱ 判断基準の遵守状況" in text
    assert "Ⅲ" in text  # 追加措置
    assert "Ⅳ" in text  # 計測データ
    assert "Ⅴ" in text  # 中長期計画進捗
    assert "Ⅵ" in text  # 提出準備チェックリスト


@pytest.mark.asyncio
async def test_form5_docx_omits_handling_section_for_chain() -> None:
    """連鎖化事業者は Ⅳ-2 荷役等時間 セクションが描画されない。"""
    result = await generate_periodic_report("sample_chain_01")
    doc = Document(result.files["docx"])
    text = "\n".join(p.text for p in doc.paragraphs)
    # 「4-2. 荷役等時間 計測実績」見出しが無いこと
    assert "4-2." not in text or "荷役等時間 計測実績" not in text


@pytest.mark.asyncio
async def test_summary_includes_compliance_and_checklist() -> None:
    """サマリーに遵守状況と提出実務チェックリストが含まれる。"""
    result = await generate_periodic_report("sample_mfg_01")
    text = build_periodic_report_summary(result)
    assert "判断基準の遵守状況" in text
    assert "計測データ" in text
    assert "中長期計画の進捗" in text
    assert "提出準備チェックリスト" in text
    assert "次のアクション" in text


@pytest.mark.asyncio
async def test_data_source_propagates_from_profile() -> None:
    """data_source は CompanyProfile から伝搬される（サンプル3社は 'sample'）。"""
    result = await generate_periodic_report("sample_mfg_01")
    assert result.data_source == "sample"


@pytest.mark.asyncio
async def test_wholesale_sample_compliance_status() -> None:
    """卸売業サンプル（中堅レベル計測）でも完成。違反項目があれば追加措置を含む。"""
    result = await generate_periodic_report("sample_wholesale_01")
    assert result.compliance.overall_status in (
        "compliant", "borderline", "violation",
    )
    # 計測データが取得されている（卸売業は荷役等あり）
    assert len(result.wait_time_rows) >= 1
    assert len(result.handling_time_rows) >= 1
