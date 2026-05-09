"""出口層 writer の統合テスト（実ファイル生成）。"""

from __future__ import annotations

from pathlib import Path

import pytest

from core.export import (
    docx_form3_writer,
    excel_writer,
    html_writer,
    markdown_writer,
    ppt_writer,
)
from core.export.citation import Citation, estat_citation


@pytest.fixture(autouse=True)
def redirect_outputs(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """outputs/ を一時ディレクトリにリダイレクト。"""
    monkeypatch.setattr("core.export.response.OUTPUTS_DIR", tmp_path)


# ---- markdown_writer ----


def test_write_macro_overview_markdown() -> None:
    path = markdown_writer.write_macro_overview(
        prefecture="東京都",
        sections=[
            ("運輸業概況", "事業所数 15,234 件"),
            ("燃料価格", "レギュラー 168 円/L"),
        ],
        citations=[Citation(label="経済センサス2021", url="https://example.com")],
        company_id="testco",
    )
    assert path.exists()
    content = path.read_text(encoding="utf-8")
    assert "東京都" in content
    assert "15,234" in content
    assert "経済センサス2021" in content


def test_write_clo_plan_markdown_includes_executive_summary() -> None:
    path = markdown_writer.write_clo_plan_draft(
        company_name="テスト株式会社",
        company_id="testco",
        plan_period="2026〜2030",
        executive_summary="経営層向け要約。年間取扱9万トン以上のため CLO 選任義務あり。",
        sections=[("Ⅱ 積載率", "本文")],
        citations=[Citation(label="A", url="https://a.example.com")],
    )
    content = path.read_text(encoding="utf-8")
    assert "エグゼクティブサマリー" in content
    assert "経営層向け要約" in content


# ---- html_writer ----


def test_write_clo_plan_html_is_self_contained() -> None:
    """生成 HTML は外部 CSS/JS 参照なしで開けること。

    社内メールに添付して誰でもダブルクリックで開ける運用が前提。
    """
    path = html_writer.write_clo_plan_draft(
        company_name="テスト株式会社",
        company_id="testco",
        plan_period="2026〜2030",
        executive_summary="営業利益+5億円・CO2削減2,400t",
        sections=[("Step 1: ベンチマーク", "業界平均比 +12%")],
        citations=[
            Citation(label="経済センサス2021", url="https://example.com/estat"),
        ],
    )
    assert path.exists()
    assert path.suffix == ".html"
    content = path.read_text(encoding="utf-8")
    # self-contained の証拠: 外部 CSS/JS の参照がない
    assert "<link rel=\"stylesheet\"" not in content
    assert "<script src=" not in content
    # 主要要素
    assert "<!DOCTYPE html>" in content
    assert "テスト株式会社" in content
    assert "計画期間" in content
    assert "営業利益+5億円" in content
    assert "Step 1: ベンチマーク" in content
    assert "業界平均比 +12%" in content
    # 出典は <a target="_blank"> でリンク化
    assert 'href="https://example.com/estat"' in content
    assert 'target="_blank"' in content


def test_write_clo_plan_html_escapes_user_content() -> None:
    """会社名や本文に含まれる HTML 特殊文字がエスケープされる。"""
    path = html_writer.write_clo_plan_draft(
        company_name="A&B<script>株式会社",
        company_id="testco",
        plan_period="2026〜2030",
        executive_summary="<img src=x onerror=alert(1)>",
        sections=[("Step", "</style><script>evil()</script>")],
        citations=[Citation(label="出典", url="https://example.com")],
    )
    content = path.read_text(encoding="utf-8")
    # 危険なタグはエスケープされていること
    assert "<script>" not in content
    assert "<img src=x" not in content
    # エスケープ済みの形では含まれていること
    assert "A&amp;B&lt;script&gt;" in content


# ---- excel_writer ----


def test_write_data_table_excel() -> None:
    from openpyxl import load_workbook

    path = excel_writer.write_data_table(
        title="運輸業事業所数 都道府県別",
        company_id="testco",
        headers=["都道府県", "事業所数"],
        rows=[["東京都", 15234], ["大阪府", 9821]],
        citations=[
            estat_citation("経済センサス2021", "0004005642"),
        ],
    )
    assert path.exists()
    wb = load_workbook(path)
    assert "Summary" in wb.sheetnames
    assert "Data" in wb.sheetnames
    assert "出典" in wb.sheetnames

    ws_summary = wb["Summary"]
    assert ws_summary.freeze_panes == "A8"
    assert len(ws_summary._charts) >= 1
    assert ws_summary.cell(row=4, column=1).value == "年間コスト削減"
    assert ws_summary.cell(row=7, column=1).value == "優先施策サマリー"

    ws = wb["Data"]
    # タイトル行・ヘッダ行・データ2行
    assert ws.cell(row=2, column=1).value == "都道府県"
    assert ws.cell(row=3, column=2).value == 15234
    assert ws.freeze_panes == "A3"
    assert "DataTable" in ws.tables

    # 出典シート
    ws_cit = wb["出典"]
    assert ws_cit.cell(row=3, column=1).value == "経済センサス2021"
    assert ws_cit.freeze_panes == "A3"
    assert "CitationTable" in ws_cit.tables


def test_write_clo_plan_excel_summary_is_human_readable() -> None:
    from openpyxl import load_workbook

    path = excel_writer.write_data_table(
        title="テクノパーツ製造株式会社 CLO 中長期計画書 数値目標・施策効果一覧",
        company_id="testco",
        headers=["指標／施策", "区分", "現状（業界平均）", "目標／効果", "出典"],
        rows=[
            ["積載率", "数値目標", "38.0%", "50.0%", "国交省"],
            [
                "積載率向上",
                "施策効果（推計）",
                "投資 約300万円",
                "CO2 約1,200.0トン/年削減・コスト 約36,000万円/年削減",
                "国交省",
            ],
            [
                "共同配送",
                "施策効果（推計）",
                "投資 約500万円",
                "CO2 約900.0トン/年削減・コスト 約30,000万円/年削減",
                "認定事例",
            ],
        ],
        citations=[
            estat_citation("経済センサス2021", "0004005642"),
        ],
        purpose="clo_plan",
    )
    wb = load_workbook(path)
    ws_summary = wb["Summary"]

    assert ws_summary.cell(row=5, column=1).value == "6.6億円"
    assert ws_summary.cell(row=8, column=1).value == "優先"
    assert ws_summary.cell(row=9, column=2).value == "積載率向上"
    assert ws_summary.cell(row=9, column=3).value == 36000
    assert ws_summary.cell(row=9, column=5).value == "<0.1年"
    assert ws_summary.cell(row=13, column=1).value == "目標達成計画"
    assert ws_summary.cell(row=15, column=1).value == "積載率"
    assert "積載率向上" in ws_summary.cell(row=15, column=4).value
    assert ws_summary.cell(row=15, column=6).value == "積載率、空車率、共同配送便数"
    assert "SummaryMeasureTable" in ws_summary.tables
    assert "TargetPlanTable" in ws_summary.tables
    assert len(ws_summary._charts) >= 1


# ---- ppt_writer ----


def test_write_clo_plan_ppt_creates_pptx() -> None:
    from pptx import Presentation

    path = ppt_writer.write_clo_plan_draft(
        company_name="テスト株式会社",
        company_id="testco",
        plan_period="2026〜2030",
        executive_summary=["年間取扱 12 万トン", "CLO 選任義務あり"],
        body_slides=[
            ppt_writer.SlideContent(
                title="Ⅱ 積載率向上",
                body_lines=["現状 70% → 目標 80%", "施策: 共同配送・ルート最適化"],
                footnote="出典: 自動車輸送統計年報",
            ),
        ],
        citations=[Citation(label="A", url="https://a.example.com")],
    )
    assert path.exists()
    prs = Presentation(str(path))
    # 表紙 + サマリー + 本文 + 出典 = 4 枚
    assert len(prs.slides) == 4
    assert prs.slide_width > prs.slide_height


def test_write_clo_executive_brief_creates_three_slide_pptx() -> None:
    """フェーズ4.1 a.3: 3 スライド経営会議要約 PPT。"""
    from pptx import Presentation

    path = ppt_writer.write_clo_executive_brief(
        company_name="テクノパーツ製造株式会社",
        company_id="sample_mfg_01",
        plan_period="2026-04 〜 2030-03",
        executive_summary=[
            "営業利益(標準): +5.2億円/年",
            "CO2削減(標準): 約 1,200トン-CO2/年",
            "炭素価格回避: 12〜30百万円/年",
            "5年累積NPV: 約 18億円",
            "会社名: テクノパーツ製造株式会社",
            "特定荷主3区分: 第一種荷主",
            "年間売上 850億円",
        ],
        actions_to_pursue=[
            "TMS 配車最適化",
            "バース予約システム導入",
            "パレット標準化",
        ],
        actions_to_skip=[
            "中継輸送（先送り）",
            "同業共同配送（パートナー候補なし）",
            "鉄道モーダルシフト（効果限定）",
        ],
        next_decision_due="2026-05-15（金）取締役会",
        citations=[Citation(label="A", url="https://a.example.com")],
    )
    assert path.exists()
    prs = Presentation(str(path))
    # 3 スライド（表紙 + 経営サマリー + やる/やらない結論）
    assert len(prs.slides) == 3
    # ファイルサイズが軽量（500 KB 未満）
    size_kb = path.stat().st_size / 1024
    assert size_kb < 500, f"PPT が {size_kb:.1f} KB と大きすぎる（目標 < 500 KB）"


def test_executive_brief_third_slide_has_decisions_and_next_due() -> None:
    """3 スライド要約の 3 枚目に「やる/やらない」と次回意思決定期日が含まれる。"""
    from pptx import Presentation

    path = ppt_writer.write_clo_executive_brief(
        company_name="X",
        company_id="x",
        plan_period="2026",
        executive_summary=["要約 1", "要約 2"],
        actions_to_pursue=["TMS 配車最適化", "バース予約", "パレット標準化"],
        actions_to_skip=["中継輸送", "同業共配", "モーダルシフト"],
        next_decision_due="2026-05-15（金）",
        citations=[],
    )
    prs = Presentation(str(path))
    slide3 = prs.slides[2]
    slide3_text = "\n".join(
        shape.text for shape in slide3.shapes if hasattr(shape, "text")
    )
    assert "やる" in slide3_text
    assert "やらない" in slide3_text
    assert "TMS 配車最適化" in slide3_text
    assert "中継輸送" in slide3_text
    assert "2026-05-15" in slide3_text


def test_ppt_executive_summary_is_second_slide() -> None:
    """エグゼクティブサマリーが必ず 2 枚目に配置される。"""
    from pptx import Presentation

    path = ppt_writer.write_clo_plan_draft(
        company_name="X",
        company_id="x",
        plan_period="2026",
        executive_summary=["要約テキスト"],
        body_slides=[],
        citations=[Citation(label="A", url="https://a.example.com")],
    )
    prs = Presentation(str(path))
    slide2 = prs.slides[1]
    slide2_text = "\n".join(shape.text for shape in slide2.shapes if hasattr(shape, "text"))
    assert "エグゼクティブサマリー" in slide2_text


def test_ppt_uses_codex_recommended_design_surface() -> None:
    """PPT が素の空白スライドではなく、デザイン要素を含む。"""
    from pptx import Presentation

    path = ppt_writer.write_clo_plan_draft(
        company_name="X",
        company_id="x",
        plan_period="2026",
        executive_summary=["推奨: 共同配送を先行", "投資額: 1億円"],
        body_slides=[],
        citations=[Citation(label="A", url="https://a.example.com")],
    )
    prs = Presentation(str(path))
    first_slide_text = "\n".join(
        shape.text for shape in prs.slides[0].shapes if hasattr(shape, "text")
    )
    assert "C L O   M I D - T E R M   P L A N   D R A F T" in first_slide_text
    assert "物流効率化法 第46条／第65条 対応" in first_slide_text


# ---- docx_form3_writer ----


def test_write_form3_docx_creates_docx() -> None:
    from docx import Document

    path = docx_form3_writer.write_form3(
        company_name="テスト株式会社",
        company_id="testco",
        main_business="その他の電子部品・デバイス製造業",
        jsic_code="3159",
        clo_name="専務執行役員 田中 太郎",
        plan_period="2026年度〜2030年度",
        sections=[
            docx_form3_writer.PlanSection(
                title="Ⅱ 積載率向上",
                shipper_class="特定第一種荷主",
                measures=[
                    docx_form3_writer.PlanMeasure(
                        measure="共同配送の推進",
                        description="同一地域への配送を集約",
                        timing="2026〜2028年度",
                    ),
                ],
            ),
        ],
        is_chain_franchise=False,
    )
    assert path.exists()
    doc = Document(str(path))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "テスト株式会社" in full_text or any(
        "テスト株式会社" in c.text for table in doc.tables for row in table.rows for c in row.cells
    )
    table_cells = [c.text for table in doc.tables for row in table.rows for c in row.cells]
    assert any("提出前レビュー用ドラフト" in text for text in table_cells)
    assert len(doc.tables) >= 3


def test_write_form3_chain_franchise_uses_section5_title() -> None:
    """連鎖化事業者は様式第3「第5条関係」のタイトル。"""
    from docx import Document

    path = docx_form3_writer.write_form3(
        company_name="チェーン本部",
        company_id="chain",
        main_business="コンビニエンスストア",
        jsic_code="5891",
        clo_name="x",
        plan_period="2026〜2030",
        sections=[],
        is_chain_franchise=True,
    )
    doc = Document(str(path))
    headings = [p.text for p in doc.paragraphs if p.style and p.style.name.startswith("Heading")]
    assert any("第5条関係" in h for h in headings)
