"""自社プロフィール入力テンプレ Excel 生成のテスト。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from core.io.template_writer import write_company_template


def test_template_creates_all_sheets(tmp_path: Path) -> None:
    """テンプレ生成で 6 シートすべてが作られる。"""
    out = write_company_template(tmp_path / "template.xlsx")
    assert out.is_file()

    wb = load_workbook(out, read_only=True)
    expected = {
        "profile",
        "sites",
        "wait_time",
        "handling_time",
        "kpi_history",
        "_jsic_reference",
    }
    assert set(wb.sheetnames) == expected


def test_profile_sheet_has_required_keys(tmp_path: Path) -> None:
    """profile シートに必須項目（company_id / jsic_code / annual_revenue_jpy）が含まれる。"""
    out = write_company_template(tmp_path / "template.xlsx")
    wb = load_workbook(out, read_only=True)
    ws = wb["profile"]
    keys = [row[0].value for row in ws.iter_rows() if row[0].value]
    must_have = (
        "company_id",
        "jsic_code",
        "annual_revenue_jpy",
        "annual_freight_tons_total",
        "operating_profit_jpy",
        "cost_of_goods_sold_jpy",
        "is_chain_franchise",
        "clo_name",
        "plan_period_start_fy",
        "plan_period_end_fy",
    )
    for key in must_have:
        assert any(
            isinstance(k, str) and k.startswith(key) for k in keys
        ), f"profile シートに {key} が存在しない（取得 keys: {keys}）"


def test_tabular_sheets_have_header_and_sample_rows(tmp_path: Path) -> None:
    """sites/wait_time/handling_time/kpi_history はヘッダ + 説明 + サンプル行を持つ。"""
    out = write_company_template(tmp_path / "template.xlsx")
    wb = load_workbook(out, read_only=True)
    for sheet_name in ("sites", "wait_time", "handling_time", "kpi_history"):
        ws = wb[sheet_name]
        # 1 行目: ヘッダ（少なくとも 4 列）
        header = [c.value for c in ws[1]]
        assert any(v for v in header), f"{sheet_name} ヘッダ行が空"
        assert len(header) >= 4

        # 3 行目: サンプル行（少なくとも 1 行）
        sample = [c.value for c in ws[3]]
        assert any(v is not None for v in sample), f"{sheet_name} サンプル行が空"


def test_jsic_reference_has_known_codes(tmp_path: Path) -> None:
    """_jsic_reference に主要業種コードが含まれる。"""
    out = write_company_template(tmp_path / "template.xlsx")
    wb = load_workbook(out, read_only=True)
    ws = wb["_jsic_reference"]
    codes = [
        ws.cell(row=row_idx, column=1).value
        for row_idx in range(2, 10)
    ]
    # サンプル3社で使っている業種が含まれる
    assert "3159" in codes  # 製造業
    assert "5514" in codes  # 卸売業
    assert "5891" in codes  # 連鎖化
