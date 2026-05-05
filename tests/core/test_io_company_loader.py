"""Excel/CSV → CompanyProfile ローダーのテスト。

正常系: テンプレ生成 → そのまま load で CompanyProfile が出る（サンプル値の整合）
異常系: 必須欠落 / 型不正 / JSIC 4桁制約違反
"""
from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from core.io.company_loader import (
    load_company_profile_from_csv,
    load_company_profile_from_excel,
)
from core.io.template_writer import write_company_template


def _replace_profile_value(path: Path, key_prefix: str, new_value: object) -> None:
    """profile シートの key_prefix で始まる項目の B 列を書き換える。"""
    wb = load_workbook(path)
    ws = wb["profile"]
    for row in ws.iter_rows(min_row=2):
        cell = row[0]
        if isinstance(cell.value, str) and cell.value.startswith(key_prefix):
            row[1].value = new_value
            break
    wb.save(path)


# --- 正常系 ---


def test_load_template_yields_valid_profile(tmp_path: Path) -> None:
    """テンプレ Excel をそのまま load すると、サンプル値で CompanyProfile が構築できる。"""
    template_path = tmp_path / "template.xlsx"
    write_company_template(template_path)

    profile = load_company_profile_from_excel(template_path)

    # テンプレのサンプル値（profile シート 3 行目以降のサンプル値）
    assert profile.company_id == "logi_shoji_01"
    assert profile.name == "ロジ商事株式会社"
    assert profile.jsic_code == "5514"
    assert profile.annual_revenue_jpy == 50_000_000_000
    assert profile.is_chain_franchise is False
    # 実データ取込なので data_source は actual 既定
    assert profile.data_source == "actual"
    # 拠点 3 件
    assert len(profile.sites) == 3
    # 計測値 3 件 / 2 件 / 履歴 3 件
    assert len(profile.wait_time_rows) == 3
    assert len(profile.handling_time_rows) == 2
    assert len(profile.kpi_history) == 3


def test_load_handles_optional_actual_fields(tmp_path: Path) -> None:
    """積載率実績/燃料消費/CO2 排出量フィールドがテンプレから取込まれる。"""
    template_path = tmp_path / "template.xlsx"
    write_company_template(template_path)
    profile = load_company_profile_from_excel(template_path)
    assert profile.load_factor_actual_pct == 42.0
    assert profile.annual_fuel_liters == 1_500_000.0
    assert profile.annual_co2_emissions_kg == 4_200_000.0


def test_load_optional_fields_blank_yields_none(tmp_path: Path) -> None:
    """積載率を空欄にすると None として読まれる。"""
    template_path = tmp_path / "template.xlsx"
    write_company_template(template_path)
    _replace_profile_value(template_path, "load_factor_actual_pct", None)

    profile = load_company_profile_from_excel(template_path)
    assert profile.load_factor_actual_pct is None


# --- 異常系 ---


def test_missing_file_raises(tmp_path: Path) -> None:
    """存在しないファイルは FileNotFoundError。"""
    with pytest.raises(FileNotFoundError):
        load_company_profile_from_excel(tmp_path / "nope.xlsx")


def test_missing_required_sheet_raises(tmp_path: Path) -> None:
    """profile シートが欠けたファイルは ValueError。"""
    path = tmp_path / "broken.xlsx"
    wb = Workbook()
    # default sheet を 'random' という名前に
    ws = wb.active
    if ws is not None:
        ws.title = "random"
    wb.save(path)
    with pytest.raises(ValueError, match="必須シート"):
        load_company_profile_from_excel(path)


def test_missing_required_field_raises(tmp_path: Path) -> None:
    """profile シートで必須フィールド（company_id）を空にすると ValueError。"""
    template_path = tmp_path / "template.xlsx"
    write_company_template(template_path)
    _replace_profile_value(template_path, "company_id", None)

    with pytest.raises(ValueError, match="company_id"):
        load_company_profile_from_excel(template_path)


def test_invalid_jsic_length_raises(tmp_path: Path) -> None:
    """jsic_code が 3 桁などだと ValueError（CompanyProfile の min_length=4）。"""
    template_path = tmp_path / "template.xlsx"
    write_company_template(template_path)
    _replace_profile_value(template_path, "jsic_code", "911")

    with pytest.raises(ValueError, match="jsic_code"):
        load_company_profile_from_excel(template_path)


def test_negative_revenue_raises(tmp_path: Path) -> None:
    """売上が負の値だと ValueError（CompanyProfile の ge=0）。"""
    template_path = tmp_path / "template.xlsx"
    write_company_template(template_path)
    _replace_profile_value(template_path, "annual_revenue_jpy", -1)

    with pytest.raises(ValueError, match="annual_revenue_jpy"):
        load_company_profile_from_excel(template_path)


def test_invalid_site_type_in_sites_sheet_raises(tmp_path: Path) -> None:
    """sites シートに不正な site_type を入れると ValueError。"""
    template_path = tmp_path / "template.xlsx"
    write_company_template(template_path)
    wb = load_workbook(template_path)
    ws = wb["sites"]
    # 3 行目（最初のサンプル行）の site_type を不正値に
    ws.cell(row=3, column=2).value = "invalid_type"
    wb.save(template_path)

    with pytest.raises(ValueError, match="sites"):
        load_company_profile_from_excel(template_path)


# --- CSV 経路 ---


def _write_csv(path: Path, rows: list[list[object]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        for row in rows:
            writer.writerow(row)


def test_load_from_csv_directory(tmp_path: Path) -> None:
    """CSV 経路でも CompanyProfile が構築できる。"""
    _write_csv(
        tmp_path / "profile.csv",
        [
            ["項目", "値"],
            ["company_id", "csv_co_01"],
            ["name", "CSV 株式会社"],
            ["jsic_code", "3159"],
            ["main_business", "電子部品製造業"],
            ["annual_revenue_jpy", 30_000_000_000],
            ["annual_freight_tons_total", 100_000],
            ["operating_profit_jpy", 1_500_000_000],
            ["cost_of_goods_sold_jpy", 22_000_000_000],
            ["is_chain_franchise", False],
            ["clo_name", "CLO 山田"],
            ["plan_period_start_fy", 2026],
            ["plan_period_end_fy", 2030],
        ],
    )
    _write_csv(
        tmp_path / "sites.csv",
        [
            ["name", "site_type", "prefecture", "annual_freight_tons"],
            ["本社", "factory", "東京都", 60_000],
            ["関西センター", "dc", "大阪府", 40_000],
        ],
    )
    # wait_time / handling_time / kpi_history は省略 → 空 list

    profile = load_company_profile_from_csv(tmp_path)

    assert profile.company_id == "csv_co_01"
    assert profile.jsic_code == "3159"
    assert len(profile.sites) == 2
    assert profile.wait_time_rows == []
    assert profile.data_source == "actual"


def test_csv_missing_directory_raises(tmp_path: Path) -> None:
    """CSV ディレクトリが存在しない場合 FileNotFoundError。"""
    with pytest.raises(FileNotFoundError):
        load_company_profile_from_csv(tmp_path / "nope")
