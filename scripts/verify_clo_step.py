"""CLO ツール群の自動検証スクリプト。

Claude Desktop を起動せずにフェーズ3.5 の全 CLO ツールをローカルで実行し、
ファイル生成 / 数値桁チェック / docx-pptx 中身検証 / 構造レポートを行う。

Usage:
    uv run python scripts/verify_clo_step.py
        # 全ツール × 全 3 サンプル一括検証

    uv run python scripts/verify_clo_step.py --tool plan_draft
        # 特定ツール × 全サンプル

    uv run python scripts/verify_clo_step.py --sample sample_mfg_01
        # 全ツール × 特定サンプル

    uv run python scripts/verify_clo_step.py --tool plan_draft --sample sample_mfg_01
        # 特定ツール × 特定サンプル

    uv run python scripts/verify_clo_step.py --json
        # 結果を JSON で出力（機械処理用）

設計方針:
    - MCP ツールではなく core/ の関数を直接呼び出し（Claude Desktop 不要）
    - e-Stat キャッシュ（core/cache.py SQLite）を活用して 2 回目以降を高速化
    - 検証は機械的な項目のみ（数値の桁・ファイル存在・docx 必須セクション・pptx 構造）
    - 「業界感覚に照らして変な値か」の判定は呼び出し元（スキル経由の Claude）に委ねる
"""
from __future__ import annotations

import argparse
import asyncio
import json
import sys
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

# プロジェクトルートをパスに追加
PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))


@dataclass
class VerifyResult:
    """単一検証の結果。"""

    tool: str
    sample: str | None
    company: str | None
    ok: bool
    issues: list[str] = field(default_factory=list)
    summary_excerpt: dict[str, Any] = field(default_factory=dict)


# ============================================================
# 個別ツール検証関数
# ============================================================


async def verify_benchmark(sample_id: str) -> VerifyResult:
    from core.analytics.benchmark import analyze_benchmark
    from core.static.sample_company import get_sample

    profile = get_sample(sample_id)
    if profile is None or not profile.sites:
        return VerifyResult("benchmark", sample_id, None, False, ["sample_not_found"])

    pref = profile.sites[0].prefecture
    issues: list[str] = []
    try:
        result = await analyze_benchmark(
            jsic_code=profile.jsic_code,
            prefecture=pref,
            revenue_jpy=profile.annual_revenue_jpy,
        )
    except Exception as e:
        return VerifyResult(
            "benchmark", sample_id, profile.name, False, [f"exception:{e!r}"],
        )

    if result.revenue_vs_avg_pct is not None and result.revenue_vs_avg_pct < 0:
        issues.append(f"revenue_vs_avg_negative:{result.revenue_vs_avg_pct}")
    if not result.citations:
        issues.append("no_citations")

    return VerifyResult(
        tool="benchmark",
        sample=sample_id,
        company=profile.name,
        ok=len(issues) == 0,
        issues=issues,
        summary_excerpt={
            "national_peer_count": result.national_peer_count,
            "regional_peer_count": result.regional_peer_count,
            "revenue_vs_avg_pct": result.revenue_vs_avg_pct,
        },
    )


async def verify_driver_gap(sample_id: str) -> VerifyResult:
    from core.analytics.driver_gap import analyze_driver_gap
    from core.static.sample_company import get_sample

    profile = get_sample(sample_id)
    if profile is None or not profile.sites:
        return VerifyResult("driver_gap", sample_id, None, False, ["sample_not_found"])

    pref = profile.sites[0].prefecture
    issues: list[str] = []
    try:
        result = await analyze_driver_gap(prefecture=pref)
    except Exception as e:
        return VerifyResult(
            "driver_gap", sample_id, profile.name, False, [f"exception:{e!r}"],
        )

    if result.national_workers is not None and (
        result.national_workers < 1_000_000 or result.national_workers > 3_000_000
    ):
        issues.append(f"national_workers_out_of_range:{result.national_workers}")

    return VerifyResult(
        tool="driver_gap",
        sample=sample_id,
        company=profile.name,
        ok=len(issues) == 0,
        issues=issues,
        summary_excerpt={
            "national_workers": result.national_workers,
            "regional_workers": result.regional_workers,
            "regional_share_pct": result.regional_share_pct,
        },
    )


def verify_eligibility(sample_id: str) -> VerifyResult:
    from core.legal.eligibility import evaluate_eligibility
    from core.static.sample_company import get_sample

    profile = get_sample(sample_id)
    if profile is None:
        return VerifyResult("eligibility", sample_id, None, False, ["sample_not_found"])

    result = evaluate_eligibility(
        jsic_code=profile.jsic_code,
        annual_freight_tons_total=profile.annual_freight_tons_total,
        annual_revenue_jpy=profile.annual_revenue_jpy,
        is_chain_franchise=profile.is_chain_franchise,
    )
    issues: list[str] = []
    if profile.is_chain_franchise and "chain_franchise" not in result.categories:
        issues.append("chain_franchise_flag_not_propagated")

    return VerifyResult(
        tool="eligibility",
        sample=sample_id,
        company=profile.name,
        ok=len(issues) == 0,
        issues=issues,
        summary_excerpt={
            "categories": list(result.categories),
            "must_appoint_clo": result.must_appoint_clo,
        },
    )


def verify_target_proposal(sample_id: str) -> VerifyResult:
    from core.analytics.target_proposal import propose_targets
    from core.static.sample_company import get_sample

    profile = get_sample(sample_id)
    if profile is None or not profile.sites:
        return VerifyResult(
            "target_proposal", sample_id, None, False, ["sample_not_found"],
        )

    issues: list[str] = []
    try:
        result = propose_targets(
            jsic_code=profile.jsic_code,
            prefecture=profile.sites[0].prefecture,
            revenue_jpy=profile.annual_revenue_jpy,
        )
    except Exception as e:
        return VerifyResult(
            "target_proposal", sample_id, profile.name, False, [f"exception:{e!r}"],
        )

    for t in result.targets:
        if t.proposed_target_unit.startswith("%") and abs(t.proposed_target) > 100:
            issues.append(f"target_pct_out_of_range:{t.name}={t.proposed_target}")

    return VerifyResult(
        tool="target_proposal",
        sample=sample_id,
        company=profile.name,
        ok=len(issues) == 0,
        issues=issues,
        summary_excerpt={
            "target_count": len(result.targets),
            "target_names": [t.name for t in result.targets],
        },
    )


def verify_simulate_measures(sample_id: str) -> VerifyResult:
    from core.analytics.simulator import simulate_measures
    from core.static.sample_company import get_sample

    profile = get_sample(sample_id)
    if profile is None or not profile.sites:
        return VerifyResult(
            "simulate_measures", sample_id, None, False, ["sample_not_found"],
        )

    issues: list[str] = []
    try:
        result = simulate_measures(
            jsic_code=profile.jsic_code,
            prefecture=profile.sites[0].prefecture,
            revenue_jpy=profile.annual_revenue_jpy,
            annual_freight_tons=profile.annual_freight_tons_total,
        )
    except Exception as e:
        return VerifyResult(
            "simulate_measures", sample_id, profile.name, False, [f"exception:{e!r}"],
        )

    for m in result.measures:
        if m.co2_reduction_kg_per_year is not None and m.co2_reduction_kg_per_year < 0:
            issues.append(f"co2_negative:{m.name}")
        if m.cost_saving_jpy_per_year is not None and m.cost_saving_jpy_per_year < 0:
            issues.append(f"cost_negative:{m.name}")
        if m.investment_jpy < 0:
            issues.append(f"investment_negative:{m.name}")

    return VerifyResult(
        tool="simulate_measures",
        sample=sample_id,
        company=profile.name,
        ok=len(issues) == 0,
        issues=issues,
        summary_excerpt={
            "measure_count": len(result.measures),
            "total_co2_reduction_t": round(result.total_co2_reduction_kg / 1000, 0),
            "total_cost_saving_oku": round(result.total_cost_saving_jpy / 1e8, 2),
            "total_investment_man": round(result.total_investment_jpy / 1e4, 0),
        },
    )


def verify_certified_match(sample_id: str) -> VerifyResult:
    from core.analytics.certified_match import find_matches
    from core.static.sample_company import get_sample

    profile = get_sample(sample_id)
    if profile is None:
        return VerifyResult(
            "certified_match", sample_id, None, False, ["sample_not_found"],
        )

    issues: list[str] = []
    try:
        result = find_matches(jsic_code=profile.jsic_code)
    except Exception as e:
        return VerifyResult(
            "certified_match", sample_id, profile.name, False, [f"exception:{e!r}"],
        )

    return VerifyResult(
        tool="certified_match",
        sample=sample_id,
        company=profile.name,
        ok=len(issues) == 0,
        issues=issues,
        summary_excerpt={
            "match_count": len(result.matches),
            "top_case_ids": [m.case.case_id for m in result.matches[:3]],
        },
    )


def verify_feasibility_check(sample_id: str) -> VerifyResult:
    from core.analytics.feasibility_check import check_feasibility
    from core.static.sample_company import get_sample

    profile = get_sample(sample_id)
    if profile is None:
        return VerifyResult(
            "feasibility_check", sample_id, None, False, ["sample_not_found"],
        )

    issues: list[str] = []
    try:
        result = check_feasibility(profile)
    except Exception as e:
        return VerifyResult(
            "feasibility_check", sample_id, profile.name, False,
            [f"exception:{e!r}"],
        )

    # 緑+黄+赤 = 全 verdict 件数
    total = result.green_count + result.yellow_count + result.red_count
    if total != len(result.measure_verdicts):
        issues.append(f"count_mismatch:total={total}_vs_verdicts={len(result.measure_verdicts)}")

    # 連鎖化事業者は荷役等が全て red
    if profile.is_chain_franchise:
        loading = [v for v in result.measure_verdicts if v.measure.form3_section == "loading"]
        for v in loading:
            if v.level != "red":
                issues.append(f"chain_loading_not_red:{v.measure.measure_id}")

    return VerifyResult(
        tool="feasibility_check",
        sample=sample_id,
        company=profile.name,
        ok=len(issues) == 0,
        issues=issues,
        summary_excerpt={
            "green": result.green_count,
            "yellow": result.yellow_count,
            "red": result.red_count,
            "kpi_load_factor_level": result.kpi_reachability[0].feasibility_level,
            "kpi_required_pp_per_year": round(
                result.kpi_reachability[0].annual_required_improvement, 1,
            ),
        },
    )


def verify_measure_eval(sample_id: str) -> VerifyResult:
    from core.analytics.measure_eval import evaluate_measures
    from core.static.sample_company import get_sample

    profile = get_sample(sample_id)
    if profile is None:
        return VerifyResult(
            "measure_eval", sample_id, None, False, ["sample_not_found"],
        )

    issues: list[str] = []
    try:
        result = evaluate_measures(profile)
    except Exception as e:
        return VerifyResult(
            "measure_eval", sample_id, profile.name, False, [f"exception:{e!r}"],
        )

    if len(result.top_recommendations) != 3:
        issues.append(f"top_n_not_3:{len(result.top_recommendations)}")
    # スコア降順チェック
    scores = [s.total_score for s in result.top_recommendations]
    if scores != sorted(scores, reverse=True):
        issues.append("top_recommendations_not_sorted")
    # スコア範囲チェック
    for s in result.all_scored:
        if not 0.0 <= s.total_score <= 100.0:
            issues.append(f"score_out_of_range:{s.measure.measure_id}={s.total_score}")

    return VerifyResult(
        tool="measure_eval",
        sample=sample_id,
        company=profile.name,
        ok=len(issues) == 0,
        issues=issues,
        summary_excerpt={
            "industry": result.industry,
            "eligible_count": result.eligible_count,
            "top_3_ids": [
                f"{s.measure.measure_id}({s.total_score:.0f})"
                for s in result.top_recommendations
            ],
            "derived_resources": result.derived_resources,
        },
    )


def verify_playbook_search(sample_id: str) -> VerifyResult:
    from core.playbook.filters import PlaybookFilter, filter_playbook
    from core.static.jsic import category_of
    from core.static.sample_company import get_sample

    profile = get_sample(sample_id)
    if profile is None:
        return VerifyResult(
            "playbook_search", sample_id, None, False, ["sample_not_found"],
        )

    industry = category_of(profile.jsic_code)

    issues: list[str] = []
    result = filter_playbook(filter_=PlaybookFilter(
        industry=industry,
        revenue_jpy=profile.annual_revenue_jpy,
        freight_tons=profile.annual_freight_tons_total,
    ))
    if not result:
        issues.append("no_measures_found_for_profile")

    return VerifyResult(
        tool="playbook_search",
        sample=sample_id,
        company=profile.name,
        ok=len(issues) == 0,
        issues=issues,
        summary_excerpt={
            "industry": industry,
            "matched_count": len(result),
            "matched_ids": [m.measure_id for m in result],
        },
    )


async def verify_plan_draft(sample_id: str) -> VerifyResult:
    """plan_draft を実行し、生成ファイル・サマリー数値・docx中身を検証。"""
    from core.analytics.plan_draft import generate_plan_draft

    issues: list[str] = []
    try:
        result = await generate_plan_draft(sample_id)
    except Exception as e:
        return VerifyResult(
            "plan_draft", sample_id, None, False, [f"exception:{e!r}"],
        )

    # ファイル存在チェック
    for kind in ("ppt", "docx", "excel", "markdown"):
        path_str = result.files.get(kind)
        if not path_str:
            issues.append(f"file_missing:{kind}")
            continue
        if not Path(path_str).exists():
            issues.append(f"file_not_on_disk:{kind}:{path_str}")

    # 数値桁チェック
    if result.simulation.total_co2_reduction_kg < 0:
        issues.append("co2_negative")
    if result.simulation.total_cost_saving_jpy > 1e12:
        issues.append(f"cost_implausibly_large:{result.simulation.total_cost_saving_jpy:.0f}")
    if result.financial.total_5y_npv_jpy < 0:
        issues.append(f"npv_negative:{result.financial.total_5y_npv_jpy:.0f}")

    # docx 中身チェック
    if "docx" in result.files:
        issues.extend(_check_form3_docx(
            Path(result.files["docx"]), result.eligibility.categories,
        ))

    # pptx 中身チェック
    if "ppt" in result.files:
        issues.extend(_check_executive_summary_pptx(Path(result.files["ppt"])))

    return VerifyResult(
        tool="plan_draft",
        sample=sample_id,
        company=result.company_name,
        ok=len(issues) == 0,
        issues=issues,
        summary_excerpt={
            "co2_reduction_t": round(result.simulation.total_co2_reduction_kg / 1000),
            "cost_saving_oku": round(result.simulation.total_cost_saving_jpy / 1e8, 2),
            "operating_profit_increase_oku": round(
                result.financial.total_operating_profit_increase_jpy / 1e8, 2,
            ),
            "npv_5y_oku": round(result.financial.total_5y_npv_jpy / 1e8, 2),
            "files": {k: str(v) for k, v in result.files.items()},
        },
    )


async def verify_submission_validate(sample_id: str) -> VerifyResult:
    from core.validators.submission_validator import validate_submission

    issues: list[str] = []
    try:
        report = await validate_submission(sample_id)
    except Exception as e:
        return VerifyResult(
            "submission_validate", sample_id, None, False, [f"exception:{e!r}"],
        )

    if not report.overall_pass:
        issues.append(f"validation_failed:{len(report.errors())}_errors")
    if report.compliance_score_pct < 90:
        issues.append(f"low_compliance_score:{report.compliance_score_pct}")

    return VerifyResult(
        tool="submission_validate",
        sample=sample_id,
        company=report.company_name,
        ok=len(issues) == 0,
        issues=issues,
        summary_excerpt={
            "overall_pass": report.overall_pass,
            "compliance_score_pct": report.compliance_score_pct,
            "error_count": len(report.errors()),
            "warning_count": len(report.warnings()),
        },
    )


async def verify_financial_translate(sample_id: str) -> VerifyResult:
    from core.analytics.financial_translate import translate_to_pnl_impact
    from core.analytics.plan_draft import generate_plan_draft
    from core.static.sample_company import get_sample

    profile = get_sample(sample_id)
    if profile is None:
        return VerifyResult(
            "financial_translate", sample_id, None, False, ["sample_not_found"],
        )

    issues: list[str] = []
    try:
        plan = await generate_plan_draft(sample_id)
        result = translate_to_pnl_impact(profile, plan.simulation)
    except Exception as e:
        return VerifyResult(
            "financial_translate", sample_id, profile.name, False,
            [f"exception:{e!r}"],
        )

    # 営業利益増 = コスト削減 の同値検証
    if abs(
        result.total_operating_profit_increase_jpy
        - plan.simulation.total_cost_saving_jpy,
    ) > 1.0:
        issues.append("op_profit_neq_cost_saving")
    if result.total_carbon_avoidance_jpy_high < result.total_carbon_avoidance_jpy_low:
        issues.append("carbon_high_lt_low")

    return VerifyResult(
        tool="financial_translate",
        sample=sample_id,
        company=profile.name,
        ok=len(issues) == 0,
        issues=issues,
        summary_excerpt={
            "operating_profit_increase_oku": round(
                result.total_operating_profit_increase_jpy / 1e8, 2,
            ),
            "margin_lift_pp": round(result.operating_profit_margin_lift_pp_total, 2),
            "carbon_low_man": round(result.total_carbon_avoidance_jpy_low / 1e4, 0),
            "carbon_high_man": round(result.total_carbon_avoidance_jpy_high / 1e4, 0),
            "npv_5y_oku": round(result.total_5y_npv_jpy / 1e8, 2),
        },
    )


# ============================================================
# docx / pptx 中身チェッカー
# ============================================================


def _check_form3_docx(path: Path, categories: list[str]) -> list[str]:
    """様式第3 docx の必須セクション + 表セル空欄をチェック。"""
    from docx import Document

    issues: list[str] = []
    if not path.exists():
        return [f"docx_not_found:{path}"]

    doc = Document(str(path))
    headings = [
        p.text for p in doc.paragraphs
        if p.style and p.style.name.startswith("Heading")
    ]

    is_chain = "chain_franchise" in (categories or [])
    required = (
        ("積載率向上", "荷待ち時間短縮")
        if is_chain
        else ("積載率向上", "荷待ち時間短縮", "荷役等時間短縮")
    )

    for section in required:
        if not any(section in h for h in headings):
            issues.append(f"docx_section_missing:{section}")

    # 表セル空欄チェック（実施措置/計画内容/実施時期）
    required_cols = ("実施措置", "計画内容", "実施時期")
    for table in doc.tables:
        if not table.rows:
            continue
        header = [c.text.strip() for c in table.rows[0].cells]
        if not all(col in header for col in required_cols):
            continue
        for row_idx, row in enumerate(table.rows[1:], start=1):
            for col_idx, cell in enumerate(row.cells[: len(required_cols)]):
                if not cell.text.strip():
                    issues.append(f"docx_empty_cell:r{row_idx}c{col_idx}")

    return issues


def _check_executive_summary_pptx(path: Path) -> list[str]:
    """PPT 2 枚目がエグゼクサマリーかつ営業利益が含まれているかチェック。"""
    from pptx import Presentation

    issues: list[str] = []
    if not path.exists():
        return [f"pptx_not_found:{path}"]

    prs = Presentation(str(path))
    if len(prs.slides) < 2:
        return [f"pptx_too_few_slides:{len(prs.slides)}"]

    slide2 = prs.slides[1]
    title = slide2.shapes.title
    if title is None or "エグゼクティブサマリー" not in title.text:
        issues.append("pptx_slide2_not_executive_summary")

    body_text = ""
    for shape in slide2.shapes:
        if shape.has_text_frame:
            body_text += shape.text_frame.text
    if "営業利益" not in body_text:
        issues.append("pptx_executive_summary_missing_operating_profit")
    if "炭素価格回避" not in body_text:
        issues.append("pptx_executive_summary_missing_carbon_avoidance")
    if "5年累積NPV" not in body_text and "5年累積 NPV" not in body_text:
        issues.append("pptx_executive_summary_missing_npv")

    return issues


# ============================================================
# オーケストレーション
# ============================================================


# (ツール名, 関数, async か)
ALL_VERIFICATIONS: list[tuple[str, Any, bool]] = [
    ("benchmark", verify_benchmark, True),
    ("driver_gap", verify_driver_gap, True),
    ("eligibility", verify_eligibility, False),
    ("target_proposal", verify_target_proposal, False),
    ("simulate_measures", verify_simulate_measures, False),
    ("certified_match", verify_certified_match, False),
    ("playbook_search", verify_playbook_search, False),
    ("measure_eval", verify_measure_eval, False),
    ("feasibility_check", verify_feasibility_check, False),
    ("plan_draft", verify_plan_draft, True),
    ("submission_validate", verify_submission_validate, True),
    ("financial_translate", verify_financial_translate, True),
]


SAMPLE_IDS = ("sample_mfg_01", "sample_wholesale_01", "sample_chain_01")


async def run_one(tool_name: str, sample_id: str) -> VerifyResult:
    for name, func, is_async in ALL_VERIFICATIONS:
        if name == tool_name:
            if is_async:
                return await func(sample_id)
            return func(sample_id)
    return VerifyResult(tool_name, sample_id, None, False, [f"unknown_tool:{tool_name}"])


async def run_all(
    tool_filter: str | None, sample_filter: str | None,
) -> list[VerifyResult]:
    results: list[VerifyResult] = []
    for tool_name, _, _ in ALL_VERIFICATIONS:
        if tool_filter and tool_name != tool_filter:
            continue
        for sample_id in SAMPLE_IDS:
            if sample_filter and sample_id != sample_filter:
                continue
            results.append(await run_one(tool_name, sample_id))
    return results


def format_human(results: list[VerifyResult]) -> str:
    lines: list[str] = ["=== CLO ツール検証結果 ===", ""]
    pass_count = sum(1 for r in results if r.ok)
    fail_count = len(results) - pass_count

    lines.append(f"合計 {len(results)} 件 / PASS {pass_count} / FAIL {fail_count}")
    lines.append("")

    for r in results:
        marker = "✅" if r.ok else "❌"
        lines.append(f"{marker} [{r.tool}] {r.sample} ({r.company or 'N/A'})")
        if r.issues:
            for i in r.issues:
                lines.append(f"    - {i}")
        if r.summary_excerpt:
            for k, v in r.summary_excerpt.items():
                if isinstance(v, dict):
                    lines.append(f"    {k}:")
                    for sk, sv in v.items():
                        lines.append(f"      {sk}: {sv}")
                else:
                    lines.append(f"    {k}: {v}")
        lines.append("")

    return "\n".join(lines)


def format_json(results: list[VerifyResult]) -> str:
    return json.dumps(
        [asdict(r) for r in results], ensure_ascii=False, indent=2,
    )


async def main() -> int:
    parser = argparse.ArgumentParser(
        description="CLO ツール群の自動検証",
    )
    parser.add_argument("--tool", help="検証対象ツール名（省略で全ツール）")
    parser.add_argument(
        "--sample",
        choices=SAMPLE_IDS,
        help="対象サンプル ID（省略で全サンプル）",
    )
    parser.add_argument("--json", action="store_true", help="JSON 形式で出力")
    args = parser.parse_args()

    results = await run_all(args.tool, args.sample)

    if args.json:
        print(format_json(results))
    else:
        print(format_human(results))

    fail_count = sum(1 for r in results if not r.ok)
    return 1 if fail_count > 0 else 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
